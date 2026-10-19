import pytest
from docx import Document as DocxDocument
from fpdf import FPDF

from app.services.chunking import chunk_pages, split_text
from app.services.parsing import ParsedPage, ParsingError, parse_file


def test_short_text_is_single_chunk():
    assert split_text("Короткий текст.", 1000, 100) == ["Короткий текст."]


def test_chunks_respect_size_and_cover_all_words():
    text = "\n\n".join(f"Абзац номер {i}. " + "слово " * 40 for i in range(20))
    chunks = split_text(text, 500, 80)
    assert len(chunks) > 3
    assert all(len(c) <= 500 + 80 + 5 for c in chunks)
    joined = " ".join(chunks)
    assert all(f"номер {i}." in joined for i in range(20))


def test_overlap_repeats_tail_of_previous_chunk():
    text = "\n".join(f"Предложение количество {i} содержит несколько слов." for i in range(30))
    chunks = split_text(text, 200, 60)
    assert len(chunks) > 2
    tail_word = chunks[0].split()[-1]
    assert tail_word in chunks[1]


def test_giant_sentence_is_hard_split():
    chunks = split_text("а" * 2500, 1000, 0)
    assert [len(c) for c in chunks] == [1000, 1000, 500]


def test_chunk_pages_keeps_page_numbers():
    pages = [ParsedPage(1, "Первая страница."), ParsedPage(2, "Вторая страница.")]
    assert [(c.page, c.text) for c in chunk_pages(pages, 1000, 100)] == [
        (1, "Первая страница."),
        (2, "Вторая страница."),
    ]


def test_parse_txt_utf8_and_cp1251(tmp_path):
    (tmp_path / "a.txt").write_text("Привет, мир", encoding="utf-8")
    (tmp_path / "b.txt").write_bytes("Привет, мир".encode("cp1251"))
    assert parse_file(tmp_path / "a.txt")[0].text == "Привет, мир"
    assert parse_file(tmp_path / "b.txt")[0].text == "Привет, мир"


def test_parse_docx_includes_tables(tmp_path):
    doc = DocxDocument()
    doc.add_paragraph("Регламент отпусков")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Срок заявления"
    table.rows[0].cells[1].text = "14 дней"
    doc.save(tmp_path / "r.docx")
    text = parse_file(tmp_path / "r.docx")[0].text
    assert "Регламент отпусков" in text and "Срок заявления | 14 дней" in text


def test_parse_pdf_returns_page_numbers(tmp_path):
    pdf = FPDF()
    for body in ("First page text", "Second page text"):
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.cell(text=body)
    pdf.output(str(tmp_path / "d.pdf"))
    pages = parse_file(tmp_path / "d.pdf")
    assert [(p.page, p.text.strip()) for p in pages] == [(1, "First page text"), (2, "Second page text")]


def test_empty_or_broken_files_raise_parsing_error(tmp_path):
    (tmp_path / "empty.txt").write_text("   \n ")
    (tmp_path / "broken.pdf").write_bytes(b"not a pdf")
    for name in ("empty.txt", "broken.pdf"):
        with pytest.raises(ParsingError):
            parse_file(tmp_path / name)
