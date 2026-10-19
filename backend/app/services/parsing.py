"""Извлечение текста из PDF / DOCX / TXT. Результат — список (номер страницы | None, текст)."""
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class ParsingError(Exception):
    pass


@dataclass
class ParsedPage:
    page: int | None  # None — формат без страниц (DOCX, TXT)
    text: str


def parse_file(path: Path) -> list[ParsedPage]:
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            pages = _parse_pdf(path)
        elif ext == ".docx":
            pages = _parse_docx(path)
        elif ext == ".txt":
            pages = _parse_txt(path)
        else:
            raise ParsingError(f"Неподдерживаемый формат: {ext}")
    except ParsingError:
        raise
    except Exception as exc:  # битый файл, шифрованный PDF и т.п.
        raise ParsingError(f"Не удалось прочитать файл: {exc}") from exc

    pages = [p for p in pages if p.text.strip()]
    if not pages:
        raise ParsingError("В файле не найден текст (возможно, это скан без текстового слоя).")
    return pages


def _parse_pdf(path: Path) -> list[ParsedPage]:
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise ParsingError("PDF зашифрован паролем.")
    return [ParsedPage(i, page.extract_text() or "") for i, page in enumerate(reader.pages, start=1)]


def _parse_docx(path: Path) -> list[ParsedPage]:
    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [ParsedPage(None, "\n".join(parts))]


def _parse_txt(path: Path) -> list[ParsedPage]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return [ParsedPage(None, raw.decode(encoding))]
        except UnicodeDecodeError:
            continue
    return [ParsedPage(None, raw.decode("utf-8", errors="replace"))]
